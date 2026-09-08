import os
import re
from typing import Any, Dict, List, Optional, Tuple
from pypdf import PdfReader
from docx import Document


# Regex patterns for contact information detection
EMAIL_REGEX = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b")
PHONE_REGEX = re.compile(
    r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}|\b(?:\+?91|0)?[6-9]\d{9}\b"
)
GITHUB_REGEX = re.compile(r"(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9_-]+)", re.IGNORECASE)
LINKEDIN_REGEX = re.compile(
    r"(?:https?://)?(?:www\.)?linkedin\.com/in/([A-Za-z0-9_%-]+)", re.IGNORECASE
)

# Section heading patterns
SECTION_HEADER_PATTERNS: List[Tuple[str, re.Pattern]] = [
    (
        "skills",
        re.compile(
            r"^(?:(?:technical|core|key|professional)\s+)?skills?(?:\s+(?:&|and)\s+(?:technologies|tools|proficiencies|competencies))?:?$",
            re.IGNORECASE,
        ),
    ),
    (
        "experience",
        re.compile(
            r"^(?:(?:work|professional|industry|relevant|employment)\s+)?(?:experience|history|employment|internships?|work\s+history):?$",
            re.IGNORECASE,
        ),
    ),
    (
        "projects",
        re.compile(
            r"^(?:(?:key|featured|academic|technical|engineering|personal)\s+)?projects?(?:\s+(?:&|and)\s+works?)?:?$",
            re.IGNORECASE,
        ),
    ),
    (
        "education",
        re.compile(
            r"^(?:academic\s+background|academics?|education(?:\s+(?:&|and)\s+(?:certifications?|courses?|training))?|qualifications?|certifications?(?:\s+(?:&|and)\s+courses?)?):?$",
            re.IGNORECASE,
        ),
    ),
    (
        "summary",
        re.compile(
            r"^(?:professional\s+)?(?:summary|profile|about\s+me|career\s+objective|objective):?$",
            re.IGNORECASE,
        ),
    ),
]


def clean_extracted_text(text: str) -> str:
    """
    Normalizes whitespace and strips non-printable characters while preserving
    line breaks and paragraph structures.
    """
    if not text:
        return ""

    # Replace null bytes
    text = text.replace("\x00", "")

    # Normalize common Unicode characters
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\u00a0", " ")  # non-breaking space
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Clean non-printable control characters except standard tabs and newlines
    cleaned_lines = []
    for line in text.split("\n"):
        # Strip trailing whitespace and non-printables
        clean_line = "".join(ch for ch in line if ch.isprintable() or ch in "\t").rstrip()
        cleaned_lines.append(clean_line)

    result = "\n".join(cleaned_lines)
    # Collapse 3+ consecutive newlines to 2
    result = re.sub(r"\n{3,}", "\n\n", result).strip()
    return result


def extract_text_from_pdf(file_path: str) -> str:
    """Extracts raw text from a PDF file using pypdf."""
    reader = PdfReader(file_path)
    extracted_pages = []

    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            extracted_pages.append(page_text)

    full_text = "\n\n".join(extracted_pages)
    return clean_extracted_text(full_text)


def extract_text_from_docx(file_path: str) -> str:
    """Extracts raw text from a DOCX file including paragraphs and tables."""
    doc = Document(file_path)
    content_parts = []

    # 1. Paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            content_parts.append(para.text.strip())

    # 2. Tables (resumes frequently format skill matrices or contact in tables)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Join cell contents with separator
                content_parts.append(" | ".join(row_cells))

    full_text = "\n".join(content_parts)
    return clean_extracted_text(full_text)


def extract_text_from_file(file_path: str, file_type: str) -> str:
    """
    Main extraction router handling PDF and DOCX documents.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Resume file not found at path: {file_path}")

    ext = file_type.lower().lstrip(".")
    if ext == "pdf":
        return extract_text_from_pdf(file_path)
    elif ext in ("docx", "doc"):
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type for extraction: {file_type}")


def extract_contact_info(text: str) -> Dict[str, Optional[str]]:
    """Extracts email, phone number, GitHub handle, and LinkedIn URL."""
    # Find email
    email_match = EMAIL_REGEX.search(text)
    email = email_match.group(0).strip() if email_match else None

    # Find phone
    phone_match = PHONE_REGEX.search(text)
    phone = phone_match.group(0).strip() if phone_match else None

    # Find GitHub handle
    github_match = GITHUB_REGEX.search(text)
    github_handle = github_match.group(1).strip() if github_match else None

    # Find LinkedIn handle/URL
    linkedin_match = LINKEDIN_REGEX.search(text)
    linkedin_url = (
        f"https://linkedin.com/in/{linkedin_match.group(1).strip()}"
        if linkedin_match
        else None
    )

    return {
        "email": email,
        "phone": phone,
        "github_handle": github_handle,
        "linkedin_url": linkedin_url,
    }


def identify_heading(line: str) -> Optional[str]:
    """
    Determines whether a line represents a canonical section heading.
    Returns the canonical section name (e.g. 'skills', 'experience') or None.
    """
    clean = line.strip().rstrip(":-_#")
    if not clean or len(clean) > 50:
        return None

    # Strip markdown symbols like ## or **
    clean_stripped = re.sub(r"^[\W_]+|[\W_]+$", "", clean).strip()

    for section_name, pattern in SECTION_HEADER_PATTERNS:
        if pattern.match(clean_stripped) or pattern.match(clean):
            return section_name

    return None


def parse_resume_sections(raw_text: str) -> Dict[str, Any]:
    """
    Segments the extracted resume text into structured sections.
    Follows Section 3.2 of docs/ai/RESUME_ANALYSIS.md.
    """
    lines = raw_text.split("\n")
    sections: Dict[str, List[str]] = {
        "header": [],
        "summary": [],
        "skills": [],
        "experience": [],
        "projects": [],
        "education": [],
        "other": [],
    }

    current_section = "header"
    detected_sections_set = set()

    for line in lines:
        heading = identify_heading(line)
        if heading:
            current_section = heading
            detected_sections_set.add(heading)
            continue

        if current_section in sections:
            sections[current_section].append(line)
        else:
            sections["other"].append(line)

    # Format extracted sections by stripping empty edges
    formatted_sections = {}
    for section_name, section_lines in sections.items():
        joined = "\n".join(section_lines).strip()
        if joined:
            formatted_sections[section_name] = joined

    # If header contains text, mark header/contact as detected
    contact_info = extract_contact_info(raw_text)
    detected_list = []
    if formatted_sections.get("header") or contact_info.get("email"):
        detected_list.append("header")

    # Order detected sections logically
    standard_order = ["summary", "skills", "experience", "projects", "education"]
    for s in standard_order:
        if s in detected_sections_set or s in formatted_sections:
            detected_list.append(s)

    # Compute text metadata
    char_count = len(raw_text)
    words = raw_text.split()
    word_count = len(words)
    line_count = len(lines)

    return {
        "sections": formatted_sections,
        "detected_sections": detected_list,
        "contact_info": contact_info,
        "metadata": {
            "char_count": char_count,
            "word_count": word_count,
            "line_count": line_count,
        },
    }


def parse_resume_file(file_path: str, file_type: str) -> Tuple[str, Dict[str, Any]]:
    """
    Convenience method extracting raw text and segmenting sections in one step.
    Returns: (raw_text, parsed_data_dict)
    """
    raw_text = extract_text_from_file(file_path, file_type)
    parsed_data = parse_resume_sections(raw_text)
    return raw_text, parsed_data
