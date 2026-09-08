import io
import os
import uuid
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.db.database import SessionLocal
from app.db.models import Resume
from app.services.resume_parser import (
    clean_extracted_text,
    extract_contact_info,
    identify_heading,
    parse_resume_sections,
    extract_text_from_docx,
    extract_text_from_file,
)

client = TestClient(app)


def generate_test_pdf_bytes(lines: list[str]) -> bytes:
    """Generates a valid PDF byte string with readable text content."""
    stream_content = "BT\n/F1 12 Tf\n50 750 Td\n15 TL\n"
    for line in lines:
        safe_line = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream_content += f"({safe_line}) '\n"
    stream_content += "ET\n"
    stream_bytes = stream_content.encode("latin1")
    stream_len = len(stream_bytes)

    pdf_str = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length {stream_len} >>
stream
{stream_content}endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000300 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
400
%%EOF"""
    return pdf_str.encode("latin1")


def generate_test_docx_bytes(
    header: str, skills: str, experience: str, projects: str, education: str
) -> bytes:
    """Generates a real DOCX file with standard resume sections and a skills table."""
    doc = Document()
    doc.add_paragraph(header)
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(skills)

    # Add a table in skills section to verify table text extraction
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Languages"
    table.rows[0].cells[1].text = "Python, TypeScript"
    table.rows[1].cells[0].text = "Databases"
    table.rows[1].cells[1].text = "PostgreSQL, Redis"

    doc.add_heading("Professional Experience", level=1)
    doc.add_paragraph(experience)
    doc.add_heading("Projects", level=1)
    doc.add_paragraph(projects)
    doc.add_heading("Education", level=1)
    doc.add_paragraph(education)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_clean_extracted_text():
    """Verify that text cleaning strips null bytes and normalizes quotes/whitespace."""
    dirty_text = "Jane\x00 Doe\r\n‘Smart Quotes’ and \u201cDouble Quotes\u201d\n\n\n\nNew Section"
    cleaned = clean_extracted_text(dirty_text)

    assert "\x00" not in cleaned
    assert "\r" not in cleaned
    assert "'Smart Quotes'" in cleaned
    assert '"Double Quotes"' in cleaned
    assert "\n\n\n" not in cleaned  # collapsed excessive newlines


def test_heading_identification_heuristics():
    """Verify that various heading casing, punctuation, and prefix styles are identified."""
    assert identify_heading("Skills") == "skills"
    assert identify_heading("TECHNICAL SKILLS:") == "skills"
    assert identify_heading("## Core Skills & Technologies") == "skills"
    assert identify_heading("Work Experience:") == "experience"
    assert identify_heading("PROFESSIONAL EXPERIENCE") == "experience"
    assert identify_heading("Projects & Works") == "projects"
    assert identify_heading("Education & Certifications") == "education"
    assert identify_heading("Career Objective") == "summary"

    # Non-headings
    assert identify_heading("Designed scalable microservices using Python and Docker.") is None
    assert identify_heading("") is None


def test_contact_info_extraction():
    """Verify extraction of email, phone, GitHub handle, and LinkedIn URL."""
    text = """
    Jane Doe
    Email: candidate.jane@gmail.com
    Phone: +91 9876543210
    Portfolio: https://github.com/janedoe-dev
    Profile: https://www.linkedin.com/in/jane-doe-profile
    """
    info = extract_contact_info(text)

    assert info["email"] == "candidate.jane@gmail.com"
    assert "9876543210" in info["phone"]
    assert info["github_handle"] == "janedoe-dev"
    assert info["linkedin_url"] == "https://linkedin.com/in/jane-doe-profile"


def test_parse_resume_sections():
    """Verify structural section segmentation from normalized raw text."""
    raw_text = """John Developer
Email: john@dev.io | GitHub: github.com/johndev

Technical Skills
Python, FastAPI, Docker, PostgreSQL, React

Work Experience
Senior Backend Engineer at Acme Corp (2021 - Present)
Built scalable REST APIs and orchestrated microservices with Docker.

Projects
SkillForge AI
Evidence-based career roadmap platform using PostgreSQL and pgvector.

Education
B.S. in Computer Engineering, Stanford University (2017 - 2021)
"""
    result = parse_resume_sections(raw_text)

    assert "header" in result["detected_sections"]
    assert "skills" in result["detected_sections"]
    assert "experience" in result["detected_sections"]
    assert "projects" in result["detected_sections"]
    assert "education" in result["detected_sections"]

    assert "Python, FastAPI" in result["sections"]["skills"]
    assert "Senior Backend Engineer" in result["sections"]["experience"]
    assert "SkillForge AI" in result["sections"]["projects"]
    assert "Stanford University" in result["sections"]["education"]

    assert result["metadata"]["word_count"] > 20
    assert result["metadata"]["char_count"] > 100


def test_docx_text_extraction():
    """Verify text extraction from real DOCX file including tables."""
    docx_bytes = generate_test_docx_bytes(
        header="Alex Smith | alex@example.com | github.com/alexsmith",
        skills="Go, Kubernetes, AWS, Terraform",
        experience="Cloud Engineer at CloudTech (2020 - 2023)",
        projects="Distributed Rate Limiter in Go",
        education="M.S. in Software Engineering, MIT",
    )

    temp_path = "/tmp/test_alex_resume.docx"
    with open(temp_path, "wb") as f:
        f.write(docx_bytes)

    try:
        text = extract_text_from_file(temp_path, "docx")
        assert "Alex Smith" in text
        assert "Go, Kubernetes" in text
        # Check table extraction
        assert "Languages | Python, TypeScript" in text
        assert "Databases | PostgreSQL, Redis" in text
        assert "Cloud Engineer" in text
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_upload_and_auto_parse_pdf():
    """Verify that uploading a PDF automatically parses raw_text and saves sections in DB."""
    pdf_lines = [
        "Jane Candidate",
        "Email: jane.candidate@example.com | Phone: +1 555-123-4567",
        "GitHub: github.com/janecandidate",
        "Technical Skills",
        "Python, PostgreSQL, Docker, Redis, FastAPI",
        "Work Experience",
        "Software Engineer at BetaCorp (2022 - 2024)",
        "Projects",
        "E-Commerce Microservices Platform",
        "Education",
        "B.S. in Computer Science",
    ]
    pdf_bytes = generate_test_pdf_bytes(pdf_lines)

    files = {
        "file": ("jane_candidate.pdf", io.BytesIO(pdf_bytes), "application/pdf")
    }
    response = client.post("/api/v1/resumes/upload", files=files)
    assert response.status_code == 201
    data = response.json()

    resume_id = data["resume_id"]
    assert "skills" in data["extracted_sections"]
    assert "experience" in data["extracted_sections"]
    assert "projects" in data["extracted_sections"]
    assert "education" in data["extracted_sections"]

    # Verify DB persistence of raw_text and parsed_data
    db = SessionLocal()
    try:
        resume = db.query(Resume).filter(Resume.id == uuid.UUID(resume_id)).first()
        assert resume is not None
        assert resume.raw_text is not None
        assert "Jane Candidate" in resume.raw_text
        assert "skills" in resume.parsed_data["sections"]
        assert "Python" in resume.parsed_data["sections"]["skills"]
        assert resume.parsed_data["contact_info"]["email"] == "jane.candidate@example.com"
        assert resume.parsed_data["contact_info"]["github_handle"] == "janecandidate"

        # Cleanup
        if os.path.exists(resume.storage_path):
            os.remove(resume.storage_path)
        db.delete(resume)
        db.commit()
    finally:
        db.close()


def test_get_resume_details_endpoint():
    """Verify that GET /api/v1/resumes/{resume_id} returns the parsed resume record."""
    pdf_lines = [
        "John Doe",
        "Email: john.doe@example.com",
        "Skills",
        "Java, Spring Boot, MySQL",
    ]
    pdf_bytes = generate_test_pdf_bytes(pdf_lines)

    files = {
        "file": ("john_doe.pdf", io.BytesIO(pdf_bytes), "application/pdf")
    }
    upload_res = client.post("/api/v1/resumes/upload", files=files)
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["resume_id"]

    # Retrieve details via GET
    detail_res = client.get(f"/api/v1/resumes/{resume_id}")
    assert detail_res.status_code == 200
    detail_data = detail_res.json()

    assert detail_data["resume_id"] == resume_id
    assert detail_data["filename"] == "john_doe.pdf"
    assert detail_data["file_type"] == "pdf"
    assert detail_data["has_raw_text"] is True
    assert "Java, Spring Boot" in detail_data["raw_text"]
    assert "skills" in detail_data["parsed_data"]["detected_sections"]

    # Clean up
    db = SessionLocal()
    try:
        resume = db.query(Resume).filter(Resume.id == uuid.UUID(resume_id)).first()
        if resume:
            if os.path.exists(resume.storage_path):
                os.remove(resume.storage_path)
            db.delete(resume)
            db.commit()
    finally:
        db.close()


def test_get_resume_not_found():
    """Verify that GET on a nonexistent UUID returns 404."""
    random_id = str(uuid.uuid4())
    res = client.get(f"/api/v1/resumes/{random_id}")
    assert res.status_code == 404
    assert "not found" in res.json()["detail"].lower()
